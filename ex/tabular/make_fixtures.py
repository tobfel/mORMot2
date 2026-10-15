# generates XLSX/CSV fixtures for the Pascal tests (openpyxl + handcrafted zip)
import openpyxl, datetime, zipfile, os, sys
from openpyxl.utils.datetime import CALENDAR_MAC_1904
out = sys.argv[1]

# 1. openpyxl workbook: types, dates, gaps, multiple + hidden sheets
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Data"
ws.append(["id", "name", "amount", "flag", "date", "datetime", "time", "pct", "formula"])
ws.append([1, "Alice", 1234.5, True, datetime.date(2024, 1, 31),
           datetime.datetime(2024, 2, 29, 13, 45, 30), datetime.time(8, 30), 0.25, None])
ws.append([2, "Bob\nZeile 2", -0.001, False, datetime.date(1900, 1, 1),
           datetime.datetime(1999, 12, 31, 23, 59, 59), datetime.time(23, 0), 1, None])
ws.append([3, "  spaces  ", 12345678901234, None, None, None, None, None, None])
ws["H2"].number_format = "0.00%"
ws["C2"].number_format = '#,##0.00 "EUR"'
ws["I2"] = "=A2*2"
ws["A6"] = 6            # gap: row 5 is missing
ws["C6"] = "sparse"     # gap: column B missing
ws2 = wb.create_sheet("Zweite Tabelle")
ws2.append(["a", "b"])
ws2.append(["x", 1.5])
ws3 = wb.create_sheet("Hidden")
ws3.sheet_state = "hidden"
ws3.append(["secret"])
wb.save(os.path.join(out, "openpyxl.xlsx"))

# 2. 1904 date system
wb = openpyxl.Workbook()
wb.epoch = CALENDAR_MAC_1904
ws = wb.active
ws.append(["d"])
ws.append([datetime.date(2020, 5, 17)])
wb.save(os.path.join(out, "date1904.xlsx"))

# 3. handcrafted: prefixed namespaces, inlineStr, rich text + rPh, str/e/b
#    types, cells without r attribute, _x000D_ escapes, absolute targets
M = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
P = 'http://schemas.openxmlformats.org/package/2006/relationships'
files = {
 '[Content_Types].xml': '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/></Types>',
 '_rels/.rels': f'<?xml version="1.0"?><Relationships xmlns="{P}"><Relationship Id="rId1" Type="{R}/officeDocument" Target="/custom/book.xml"/></Relationships>',
 'custom/book.xml': f'<?xml version="1.0"?><x:workbook xmlns:x="{M}" xmlns:r="{R}"><x:sheets><x:sheet name="Hand &amp; Made" sheetId="7" r:id="rIdS"/></x:sheets></x:workbook>',
 'custom/_rels/book.xml.rels': f'<?xml version="1.0"?><Relationships xmlns="{P}"><Relationship Id="rIdS" Type="{R}/worksheet" Target="sheets/s.xml"/><Relationship Id="rIdT" Type="{R}/sharedStrings" Target="../custom/strings.xml"/></Relationships>',
 'custom/strings.xml': f'''<?xml version="1.0"?><x:sst xmlns:x="{M}" count="2" uniqueCount="2">
  <x:si><x:r><x:rPr><x:b/></x:rPr><x:t>Bold</x:t></x:r><x:r><x:t xml:space="preserve"> and plain</x:t></x:r><x:rPh sb="0" eb="1"><x:t>PHONETIC</x:t></x:rPh></x:si>
  <x:si><x:t>line1_x000D_
line2 _x005F_x0041_ &lt;tag&gt;</x:t></x:si>
</x:sst>''',
 'custom/sheets/s.xml': f'''<?xml version="1.0"?>
<x:worksheet xmlns:x="{M}">
 <x:sheetData>
  <x:row>
   <x:c t="inlineStr"><x:is><x:t>col A</x:t></x:is></x:c>
   <x:c t="inlineStr"><x:is><x:t>col B</x:t></x:is></x:c>
   <x:c t="inlineStr"><x:is><x:t>col C</x:t></x:is></x:c>
   <x:c t="inlineStr"><x:is><x:t>col A</x:t></x:is></x:c>
  </x:row>
  <x:row r="3">
   <x:c r="A3" t="s"><x:v>0</x:v></x:c>
   <x:c r="B3" t="s"><x:v>1</x:v></x:c>
   <x:c r="D3" t="str"><x:f>CONCAT("a","b")</x:f><x:v>ab</x:v></x:c>
  </x:row>
  <x:row r="4">
   <x:c r="A4" t="b"><x:v>1</x:v></x:c>
   <x:c r="B4" t="e"><x:v>#DIV/0!</x:v></x:c>
   <x:c r="C4"><x:v>1.5E-3</x:v></x:c>
   <x:c r="D4"><x:v>-42</x:v></x:c>
   <x:c r="E4"><x:v>7</x:v></x:c>
  </x:row>
 </x:sheetData>
</x:worksheet>''',
}
with zipfile.ZipFile(os.path.join(out, "handmade.xlsx"), "w", zipfile.ZIP_DEFLATED) as z:
    for n, c in files.items():
        z.writestr(n, c)

# 4. CSV fixtures written by Python's csv module (reference RFC writer)
import csv, io
rows = [["id", "text", "num"], ["1", 'He said "hi"', "1.5"], ["2", "multi\r\nline", "-3"],
        ["3", "a,b;c", ""], ["4", "", "007"]]
with open(os.path.join(out, "python.csv"), "w", newline="", encoding="utf-8") as f:
    csv.writer(f).writerows(rows)
with open(os.path.join(out, "excel_de.csv"), "w", newline="", encoding="cp1252") as f:
    f.write("sep=;\r\nName;Betrag;Datum\r\nMüller;1,5;31.01.2024\r\nGrüße;-2,25;01.02.2024\r\n")
with open(os.path.join(out, "utf16.csv"), "w", newline="", encoding="utf-16") as f:
    f.write("a\tb\r\nÄ\t€\r\n")
print("fixtures done")
